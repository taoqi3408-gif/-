"""
STDF 分析结果导出 Excel（多 sheet）
"""
import os, time, datetime
from collections import defaultdict, Counter
import numpy as np
import pandas as pd
from pystdf.IO import Parser
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import ColorScaleRule, CellIsRule

STDF_PATH = "/home/user/stdf_analysis/test.stdf"
OUT_XLSX  = "/home/user/stdf_analysis/STDF_Analysis_X15274_CV9480-A11.xlsx"

# ===== 数据容器 =====
mir_info = {}
mrr_info = {}
sdr_info = []
hbr_defs = {}
sbr_defs = {}

parts = []            # 每颗 die
ptr_per_test = defaultdict(lambda: {
    "tnum": None, "tname": "", "units": "", "ll": None, "ul": None,
    "n": 0, "fail": 0, "vals": [],
    "site_n": Counter(), "site_fail": Counter(),
})
fails_by_part = defaultdict(list)
part_meta = {}                       # idx -> {site, hbin, sbin, pass, part_id}
site_to_current_part = {}
part_idx_counter = 0

class Sink:
    def after_begin(self, ds): pass
    def after_complete(self, ds): pass
    def after_cancel(self, ds, exc): pass
    def after_send(self, ds, data):
        global part_idx_counter
        rec_type, fields = data
        name = rec_type.__class__.__name__
        d = dict(zip(rec_type.fieldNames, fields))

        if name == "Mir":
            mir_info.update(d)
        elif name == "Mrr":
            mrr_info.update(d)
        elif name == "Sdr":
            sdr_info.append(d)
        elif name == "Hbr":
            hbr_defs[d.get("HBIN_NUM")] = d
        elif name == "Sbr":
            sbr_defs[d.get("SBIN_NUM")] = d
        elif name == "Pir":
            key = (d.get("HEAD_NUM"), d.get("SITE_NUM"))
            site_to_current_part[key] = part_idx_counter
            part_meta[part_idx_counter] = {"site": d.get("SITE_NUM")}
            part_idx_counter += 1
        elif name == "Prr":
            head, site = d.get("HEAD_NUM"), d.get("SITE_NUM")
            pi = site_to_current_part.pop((head, site), None)
            if pi is not None:
                pf = d.get("PART_FLG", 0)
                meta = part_meta[pi]
                meta.update({
                    "part_idx": pi,
                    "part_id": d.get("PART_ID"),
                    "hbin": d.get("HARD_BIN"),
                    "sbin": d.get("SOFT_BIN"),
                    "pass": (pf & 0x08) == 0,
                    "x": d.get("X_COORD"),
                    "y": d.get("Y_COORD"),
                    "test_time_ms": d.get("TEST_T"),
                    "num_test": d.get("NUM_TEST"),
                })
        elif name == "Ptr":
            tnum = d.get("TEST_NUM")
            tname = d.get("TEST_TXT", "")
            key = (tnum, tname)
            head, site = d.get("HEAD_NUM"), d.get("SITE_NUM")
            pi = site_to_current_part.get((head, site))
            res = d.get("RESULT")
            tflg = d.get("TEST_FLG", 0)
            pflg = d.get("PARM_FLG", 0)
            is_fail = bool(tflg & 0xC0) or bool(pflg & 0x08)
            s = ptr_per_test[key]
            s["tnum"], s["tname"] = tnum, tname
            s["n"] += 1
            s["site_n"][site] += 1
            if is_fail:
                s["fail"] += 1
                s["site_fail"][site] += 1
                if pi is not None:
                    fails_by_part[pi].append((tnum, tname))
            if res is not None and not (tflg & 0x40):
                s["vals"].append(res)
            opt = d.get("OPT_FLAG") or 0
            if not (opt & 0x40) and d.get("LO_LIMIT") is not None:
                s["ll"] = d.get("LO_LIMIT")
            if not (opt & 0x80) and d.get("HI_LIMIT") is not None:
                s["ul"] = d.get("HI_LIMIT")
            if d.get("UNITS"):
                s["units"] = d.get("UNITS")
        elif name == "Ftr":
            tnum = d.get("TEST_NUM")
            tname = d.get("TEST_TXT", "")
            key = ("FTR", tnum, tname)
            head, site = d.get("HEAD_NUM"), d.get("SITE_NUM")
            pi = site_to_current_part.get((head, site))
            tflg = d.get("TEST_FLG", 0)
            is_fail = bool(tflg & 0xC0)
            s = ptr_per_test[key]
            s["tnum"], s["tname"] = tnum, f"[FTR]{tname}"
            s["n"] += 1
            s["site_n"][site] += 1
            if is_fail:
                s["fail"] += 1
                s["site_fail"][site] += 1
                if pi is not None:
                    fails_by_part[pi].append((tnum, tname))

print("Parsing STDF...", flush=True)
t0 = time.time()
p = Parser(inp=open(STDF_PATH, "rb"))
p.addSink(Sink())
p.parse()
print(f"Parse done in {time.time()-t0:.1f}s\n", flush=True)

# ===== 整理数据 =====
df_parts = pd.DataFrame(part_meta.values())
df_parts = df_parts.dropna(subset=["part_idx"])
df_parts["part_idx"] = df_parts["part_idx"].astype(int)
df_parts = df_parts.sort_values("part_idx").reset_index(drop=True)

# HBin name 注入
df_parts["hbin_name"] = df_parts["hbin"].map(lambda h: hbr_defs.get(h, {}).get("HBIN_NAM", ""))
df_parts["sbin_name"] = df_parts["sbin"].map(lambda s: sbr_defs.get(s, {}).get("SBIN_NAM", ""))

# 失效 die 的失效 test 拼接
def fail_tests_str(idx):
    fl = fails_by_part.get(idx, [])
    if not fl: return ""
    # 同 tnum 合并
    by_num = defaultdict(list)
    for tn, tname in fl:
        by_num[tn].append(tname)
    parts_s = []
    for tn, names in sorted(by_num.items()):
        if len(names) <= 3:
            parts_s.append(f"T{tn}[{'; '.join(names)}]")
        else:
            parts_s.append(f"T{tn}×{len(names)}[{names[0]} ...]")
    return " | ".join(parts_s)[:2000]

df_parts["fail_tests"] = df_parts["part_idx"].map(fail_tests_str)
df_parts["num_fail_tests"] = df_parts["part_idx"].map(lambda i: len(fails_by_part.get(i, [])))

# 测试统计表
def cp_cpk(vals, ll, ul):
    arr = np.asarray(vals, dtype=float)
    arr = arr[np.isfinite(arr)]
    if len(arr) < 5:
        return [None]*7
    mean = float(arr.mean())
    std  = float(arr.std(ddof=1))
    mn, mx = float(arr.min()), float(arr.max())
    cp = cpk = None
    if std > 0 and ll is not None and ul is not None:
        cp  = (ul - ll) / (6*std)
        cpk = min((mean-ll)/(3*std), (ul-mean)/(3*std))
    elif std > 0:
        if ll is not None: cpk = (mean-ll)/(3*std)
        if ul is not None: cpk = (ul-mean)/(3*std)
    return mean, std, mn, mx, cp, cpk, len(arr)

rows = []
for key, s in ptr_per_test.items():
    mean, std, mn, mx, cp, cpk, n_val = cp_cpk(s["vals"], s["ll"], s["ul"])
    rec = "FTR" if isinstance(key, tuple) and len(key)==3 and key[0]=="FTR" else "PTR"
    fr = s["fail"]/s["n"]*100 if s["n"] else 0
    rows.append({
        "RecType": rec,
        "TestNum": s["tnum"],
        "TestName": s["tname"],
        "Units": s["units"],
        "LL": s["ll"],
        "UL": s["ul"],
        "N_Exec": s["n"],
        "N_Fail": s["fail"],
        "FailRate(%)": round(fr, 3),
        "N_Valid": n_val,
        "Mean": mean,
        "Std": std,
        "Min": mn,
        "Max": mx,
        "Cp": round(cp, 3) if cp is not None else None,
        "Cpk": round(cpk, 3) if cpk is not None else None,
        "S1_n":    s["site_n"].get(1,0),
        "S1_fail": s["site_fail"].get(1,0),
        "S2_n":    s["site_n"].get(2,0),
        "S2_fail": s["site_fail"].get(2,0),
        "S3_n":    s["site_n"].get(3,0),
        "S3_fail": s["site_fail"].get(3,0),
        "S4_n":    s["site_n"].get(4,0),
        "S4_fail": s["site_fail"].get(4,0),
    })
df_tests = pd.DataFrame(rows)
for i in (1,2,3,4):
    df_tests[f"S{i}_FailRate(%)"] = np.where(
        df_tests[f"S{i}_n"]>0,
        (df_tests[f"S{i}_fail"]/df_tests[f"S{i}_n"]*100).round(3),
        np.nan)

# ===== 写 Excel =====
print(f"Writing Excel -> {OUT_XLSX}\n")

# --- Sheet: 00_Summary ---
def fmt_t(v):
    try: return datetime.datetime.utcfromtimestamp(v).strftime("%Y-%m-%d %H:%M:%S UTC")
    except: return v

total = len(df_parts)
pass_n = int(df_parts["pass"].sum())
fail_n = total - pass_n
yield_pct = pass_n/total*100 if total else 0
summary_rows = [
    ("== Lot 信息 ==", ""),
    ("Lot ID",        mir_info.get("LOT_ID")),
    ("Sub-lot ID",    mir_info.get("SBLOT_ID")),
    ("Part Type",     mir_info.get("PART_TYP")),
    ("Job Name",      mir_info.get("JOB_NAM")),
    ("Test Code",     mir_info.get("TEST_COD")),
    ("Flow ID",       mir_info.get("FLOW_ID")),
    ("Operator",      mir_info.get("OPER_NAM")),
    ("Tester Type",   mir_info.get("TSTR_TYP")),
    ("Tester Node",   mir_info.get("NODE_NAM")),
    ("SmartTest Ver", mir_info.get("EXEC_VER")),
    ("Setup Time",    fmt_t(mir_info.get("SETUP_T"))),
    ("Start Time",    fmt_t(mir_info.get("START_T"))),
    ("Finish Time",   fmt_t(mrr_info.get("FINISH_T"))),
    ("", ""),
    ("== 总体良率 ==", ""),
    ("Total Parts",   total),
    ("Pass",          f"{pass_n} ({pass_n/total*100:.2f}%)" if total else 0),
    ("Fail",          f"{fail_n} ({fail_n/total*100:.2f}%)" if total else 0),
    ("Yield",         f"{yield_pct:.2f}%"),
    ("", ""),
    ("== 测试项统计 ==", ""),
    ("Unique Tests",  len(ptr_per_test)),
    ("Site Count",    4),
    ("Parts per Site (avg)", total // 4 if total else 0),
]
df_summary = pd.DataFrame(summary_rows, columns=["Item", "Value"])

# --- Sheet: 02_Site_Yield ---
site_grp = df_parts.groupby("site").agg(
    Total=("pass","size"),
    Pass=("pass","sum"),
).reset_index()
site_grp["Fail"] = site_grp["Total"] - site_grp["Pass"]
site_grp["Yield(%)"] = (site_grp["Pass"]/site_grp["Total"]*100).round(3)
# Site x HBin
site_hbin = df_parts.groupby(["site","hbin"]).size().unstack(fill_value=0)
site_hbin.columns = [f"HB{int(c)}_{hbr_defs.get(c,{}).get('HBIN_NAM','')}" for c in site_hbin.columns]
df_site_yield = site_grp.merge(site_hbin.reset_index(), on="site")

# --- Sheet: 03_HBin ---
hbin_count = df_parts["hbin"].value_counts().reset_index()
hbin_count.columns = ["HBin", "Count"]
hbin_count["Ratio(%)"] = (hbin_count["Count"]/total*100).round(3)
hbin_count["P/F"] = hbin_count["HBin"].map(lambda h: hbr_defs.get(h,{}).get("HBIN_PF",""))
hbin_count["Name"] = hbin_count["HBin"].map(lambda h: hbr_defs.get(h,{}).get("HBIN_NAM",""))
hbin_count = hbin_count[["HBin","Name","P/F","Count","Ratio(%)"]]

# --- Sheet: 04_SBin ---
sbin_count = df_parts["sbin"].value_counts().reset_index()
sbin_count.columns = ["SBin", "Count"]
sbin_count["Ratio(%)"] = (sbin_count["Count"]/total*100).round(3)
sbin_count["P/F"] = sbin_count["SBin"].map(lambda s: sbr_defs.get(s,{}).get("SBIN_PF",""))
sbin_count["Name"] = sbin_count["SBin"].map(lambda s: sbr_defs.get(s,{}).get("SBIN_NAM",""))
sbin_count = sbin_count[["SBin","Name","P/F","Count","Ratio(%)"]]

# --- Sheet: 05_Top_Fail_Tests ---
df_top_fail = df_tests[df_tests["N_Fail"]>0].sort_values("N_Fail", ascending=False).head(100).reset_index(drop=True)

# --- Sheet: 06_All_Tests_Stats ---
df_all_tests = df_tests.sort_values(["N_Fail","TestNum"], ascending=[False, True]).reset_index(drop=True)

# --- Sheet: 07_Site_x_Test_FailRate ---
site_x_test = df_tests[df_tests["N_Fail"]>0][
    ["TestNum","TestName","N_Exec","N_Fail","FailRate(%)",
     "S1_FailRate(%)","S2_FailRate(%)","S3_FailRate(%)","S4_FailRate(%)"]
].sort_values("N_Fail", ascending=False).head(200).reset_index(drop=True)

# --- Sheet: 08_All_Parts ---
df_all_parts = df_parts[["part_idx","part_id","site","hbin","hbin_name","sbin","sbin_name",
                          "pass","x","y","test_time_ms","num_test","num_fail_tests"]].copy()

# --- Sheet: 09_Failed_Parts ---
df_failed = df_parts[~df_parts["pass"]][
    ["part_idx","part_id","site","hbin","hbin_name","sbin","sbin_name",
     "num_fail_tests","fail_tests"]].copy().reset_index(drop=True)

# --- Sheet: 10_SBin65535 ---
sb_65535_parts = df_parts[df_parts["sbin"]==65535]
sb65535_test_cnt = Counter()
for pi in sb_65535_parts["part_idx"]:
    for tn, tname in fails_by_part.get(pi, []):
        sb65535_test_cnt[(tn, tname)] += 1
df_sb65535 = pd.DataFrame(
    [{"TestNum":k[0], "TestName":k[1], "FailCount":v}
     for k,v in sb65535_test_cnt.most_common()]
)

# --- Sheet: 11_Action_Items ---
actions = [
    ("P0", "Site 4 DFT_fail 异常 (6.95% vs 平均 1.5%)",
     "清洁 S4 探针/socket，重测一批做相关性验证；检查 S4 在 DFT pin map / 电源完整性",
     "Site 4 HBin 4 占比为 Site 2 的 18 倍"),
    ("P0", "TX_RTERM_TRIM (T54000) Cp=0.07, Cpk=0.01",
     "复核 95-105 Ω 规格合理性；检查 trim 算法收敛；研究 Site 3 完全 0% fail 的原因（路径差异？）",
     "Std=24 远大于 spec width=10；16 个 TX pin 完全同步失效"),
    ("P1", "RX_GOOD_PHASE0 (T575) Site 3 失效率 10.5%",
     "排查 Site 3 RX recovery / clock 链；与其他 site 做硬件对照",
     "其他 site 仅 3.6~7.9%"),
    ("P1", "RX_RESISTANCE (T53000) S2/S4 = 0% 失效",
     "拉 S1/S3 失效 die 到 S2/S4 重测；检查 S2/S4 RX 端 reference 电阻测量链路",
     "S1=4.7%, S3=5.8%, S2/S4=0% — 站间分布二极化"),
    ("P1", "RX_GOOD_PHASE0/1 整体 Cpk≈0.55", "复核 RX recovery 设计裕度",
     "中心 11/12 vs LL=1, UL=70 — 偏向下限"),
    ("P2", "SBin 65535 缺失定义 (169 颗)",
     "为 TX_RTERM_TRIM fail 路径新增专属 SBin (如 SBin 540 = TX_RTERM_fail)",
     "168/169 颗最终落 HBin 5"),
    ("P2", "测试时间优化",
     "排序 Top 10 耗时 test，看能否并行/裁剪",
     f"Mean ~ {df_parts['test_time_ms'].mean()/1000:.1f}s/颗" if df_parts["test_time_ms"].notna().any() else ""),
]
df_actions = pd.DataFrame(actions, columns=["Priority","Issue","Recommended Action","Evidence"])

# ===== 写入 Excel + 格式化 =====
with pd.ExcelWriter(OUT_XLSX, engine="openpyxl") as writer:
    df_summary.to_excel    (writer, sheet_name="00_Summary",          index=False)
    df_actions.to_excel    (writer, sheet_name="01_Action_Items",     index=False)
    df_site_yield.to_excel (writer, sheet_name="02_Site_Yield",       index=False)
    hbin_count.to_excel    (writer, sheet_name="03_HBin_Pareto",      index=False)
    sbin_count.to_excel    (writer, sheet_name="04_SBin_Pareto",      index=False)
    df_top_fail.to_excel   (writer, sheet_name="05_Top_Fail_Tests",   index=False)
    df_all_tests.to_excel  (writer, sheet_name="06_All_Tests_Stats",  index=False)
    site_x_test.to_excel   (writer, sheet_name="07_Site_x_Test",      index=False)
    df_all_parts.to_excel  (writer, sheet_name="08_All_Parts",        index=False)
    df_failed.to_excel     (writer, sheet_name="09_Failed_Parts",     index=False)
    df_sb65535.to_excel    (writer, sheet_name="10_SBin65535",        index=False)

    # 格式化
    wb = writer.book
    header_fill = PatternFill("solid", fgColor="305496")
    header_font = Font(color="FFFFFF", bold=True)
    thin = Side(border_style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")
    p0_fill = PatternFill("solid", fgColor="FFC7CE")
    p1_fill = PatternFill("solid", fgColor="FFEB9C")
    p2_fill = PatternFill("solid", fgColor="DDEBF7")

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        # 表头格式
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
        # 冻结首行
        ws.freeze_panes = "A2"
        # 自动列宽
        for col_idx, col in enumerate(ws.columns, 1):
            max_len = 0
            for cell in col:
                v = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(v), 60))
            ws.column_dimensions[get_column_letter(col_idx)].width = max_len + 2

    # 03 HBin / 04 SBin：根据 P/F 着色
    for sh in ("03_HBin_Pareto","04_SBin_Pareto"):
        ws = wb[sh]
        pf_col = None
        for i, cell in enumerate(ws[1], 1):
            if cell.value == "P/F":
                pf_col = i; break
        if pf_col:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                v = row[pf_col-1].value
                fill = pass_fill if v=="P" else (fail_fill if v=="F" else None)
                if fill:
                    for c in row: c.fill = fill

    # 01 Action_Items：根据 Priority 着色
    ws = wb["01_Action_Items"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        pv = row[0].value
        fill = p0_fill if pv=="P0" else (p1_fill if pv=="P1" else (p2_fill if pv=="P2" else None))
        if fill:
            for c in row: c.fill = fill

    # 05/06/07: FailRate(%) 列做颜色渐变
    for sh in ("05_Top_Fail_Tests","06_All_Tests_Stats","07_Site_x_Test"):
        ws = wb[sh]
        for i, cell in enumerate(ws[1], 1):
            if cell.value and "FailRate" in str(cell.value):
                col = get_column_letter(i)
                rng = f"{col}2:{col}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                   mid_type="num", mid_value=5, mid_color="FFEB9C",
                                   end_type="num", end_value=15, end_color="FF6B6B"))
            if cell.value in ("Cp","Cpk"):
                col = get_column_letter(i)
                rng = f"{col}2:{col}{ws.max_row}"
                ws.conditional_formatting.add(rng,
                    ColorScaleRule(start_type="num", start_value=0, start_color="FF6B6B",
                                   mid_type="num", mid_value=1.33, mid_color="FFEB9C",
                                   end_type="num", end_value=2.0, end_color="63BE7B"))

    # 09 Failed_Parts: 整行淡红
    ws = wb["09_Failed_Parts"]
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row: c.fill = fail_fill

print("Done.")
print(f"File: {OUT_XLSX}")
print(f"Size: {os.path.getsize(OUT_XLSX)/1024/1024:.2f} MB")
